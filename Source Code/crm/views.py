import csv
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required, user_passes_test
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q
from django.db.models.deletion import ProtectedError
from django.http import HttpResponse, Http404
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from .models import Customer, Product, Order, Payment, FollowUp, SalesTarget, AuditLog, JourneyEvent
from .forms import CustomerForm, ProductForm, OrderForm, PaymentForm, FollowUpForm, TargetForm, Items

CONFIG={
 'customers':(Customer,CustomerForm,'Customers','Customer'),
 'orders':(Order,OrderForm,'Orders & sales','Order'),
 'followups':(FollowUp,FollowUpForm,'Follow-ups','Follow-up'),
 'products':(Product,ProductForm,'Product catalogue','Product'),
 'payments':(Payment,PaymentForm,'Payments','Payment'),
 'targets':(SalesTarget,TargetForm,'Monthly targets','Target'),
}
REVENUE_STATUSES=['confirmed','processing','completed']

def audit(user,action,obj):
    AuditLog.objects.create(actor=user,action=action,entity=obj._meta.model_name,object_id=obj.pk,label=str(obj)[:250])

def scoped(user,kind):
    if kind not in CONFIG: raise Http404
    if kind=='targets' and not user.is_staff: raise PermissionDenied
    model=CONFIG[kind][0]
    qs=model.objects.all()
    if kind=='customers': qs=qs.select_related('owner')
    elif kind in ('orders','followups'): qs=qs.select_related('customer','owner')
    elif kind=='payments': qs=qs.select_related('order','order__customer')
    elif kind=='targets': qs=qs.select_related('user')
    if not user.is_staff:
        if kind in ('customers','orders','followups'): qs=qs.filter(owner=user)
        if kind=='payments': qs=qs.filter(order__owner=user)
    if kind=='orders': qs=qs.prefetch_related('items','payments')
    return qs

def filtered(request,kind):
    qs=scoped(request.user,kind)
    q=request.GET.get('q','').strip()
    if q:
        fields={'customers':['name','company','email','phone'],'orders':['customer__name','notes'],'followups':['subject','customer__name'],'products':['name','sku'],'payments':['reference','order__customer__name'],'targets':['user__username']}[kind]
        clause=Q()
        for f in fields: clause|=Q(**{f+'__icontains':q})
        if kind=='orders' and q.removeprefix('SO-').isdigit(): clause|=Q(pk=int(q.removeprefix('SO-')))
        qs=qs.filter(clause)
    field='stage' if kind=='customers' else ('status' if kind in ('orders','followups') else None)
    if field and request.GET.get('status'): qs=qs.filter(**{field:request.GET['status']})
    if request.user.is_staff and request.GET.get('owner','').isdigit() and kind in ('customers','orders','followups'):
        qs=qs.filter(owner_id=request.GET['owner'])
    if request.GET.get('start') and request.GET.get('end'):
        try:
            if date.fromisoformat(request.GET['start'])>date.fromisoformat(request.GET['end']): raise ValueError
        except ValueError: raise ValueError('Choose a valid date range with start before end.')
    datefield={'customers':'created_at__date','orders':'order_date','followups':'due_at__date','payments':'paid_on','targets':'month'}.get(kind)
    for param,lookup in [('start','gte'),('end','lte')]:
        value=request.GET.get(param)
        if value and datefield:
            try: parsed=date.fromisoformat(value)
            except ValueError: raise ValueError('Use valid dates in YYYY-MM-DD format.')
            qs=qs.filter(**{datefield+'__'+lookup:parsed})
    if request.GET.get('overdue')=='1' and kind=='followups': qs=qs.filter(status='open',due_at__lt=timezone.now())
    return qs

@login_required
def listing(request,kind):
    if kind not in CONFIG: raise Http404
    try: qs=filtered(request,kind)
    except ValueError as exc: return HttpResponse(str(exc),status=400)
    model,_,title,singular=CONFIG[kind]
    choices=model._meta.get_field('stage' if kind=='customers' else 'status').choices if kind in ('customers','orders','followups') else []
    params=request.GET.copy(); params.pop('page',None)
    return render(request,'crm/list.html',{'kind':kind,'title':title,'singular':singular,'page':Paginator(qs,20).get_page(request.GET.get('page')),'choices':choices,'query':params.urlencode(),'team':get_user_model().objects.filter(is_active=True),'can_edit':kind!='products' or request.user.is_staff})

@login_required
@transaction.atomic
def edit(request,kind,pk=None):
    if kind not in CONFIG: raise Http404
    model,formclass,title,singular=CONFIG[kind]
    if kind in ('targets','products') and not request.user.is_staff: raise PermissionDenied
    obj=get_object_or_404(scoped(request.user,kind).select_for_update(),pk=pk) if pk else model()
    previous_stage=obj.stage if kind=='customers' and pk else ''
    old_owner=obj.owner_id if kind=='customers' and pk else None
    old_updated=str(obj.updated_at.isoformat()) if pk and hasattr(obj,'updated_at') else ''
    if not pk and 'customer' in [f.name for f in model._meta.fields]:
        cid=request.GET.get('customer')
        if cid and cid.isdigit():
            c=get_object_or_404(scoped(request.user,'customers'),pk=cid)
            obj.customer=c; obj.owner=c.owner
    if hasattr(obj,'owner_id') and not obj.owner_id: obj.owner=request.user
    form=formclass(request.POST or None,instance=obj,user=request.user)
    items=Items(request.POST or None,instance=obj,prefix='items') if kind=='orders' else None
    valid=form.is_valid() if request.method=='POST' else False
    items_valid=items.is_valid() if items is not None and request.method=='POST' else True
    if request.method=='POST' and old_updated and request.POST.get('revision')!=old_updated:
        form.add_error(None,'This record changed after you opened it. Reload this page before saving.'); valid=False
    if valid and items_valid:
        obj=form.save(commit=False)
        if hasattr(obj,'owner_id') and not request.user.is_staff: obj.owner=request.user
        if kind=='orders':
            rows=[f.cleaned_data for f in items.forms if f.cleaned_data and not f.cleaned_data.get('DELETE')]
            subtotal=sum((d['quantity']*d['unit_price'] for d in rows),Decimal(0))
            if obj.discount>subtotal: form.add_error('discount','Discount cannot exceed the subtotal.')
            total=((subtotal-obj.discount)*(1+obj.tax_percent/100)).quantize(Decimal('.01'))
            if pk and obj.paid>total: form.add_error(None,'Order total cannot be less than recorded payments.')
            if pk and obj.paid and obj.status in ['draft','cancelled']: form.add_error('status','Remove payments before reverting or cancelling this order.')
        if kind=='payments':
            order=Order.objects.select_for_update().get(pk=obj.order_id)
            other_paid=sum((p.amount for p in order.payments.exclude(pk=obj.pk)),Decimal(0))
            if order.status not in REVENUE_STATUSES: form.add_error('order','Only confirmed, processing or completed orders accept payments.')
            if obj.amount+other_paid>order.total: form.add_error('amount','Payment would exceed the order balance.')
        if not form.errors:
            if kind=='followups': obj.completed_at=(obj.completed_at or timezone.now()) if obj.status=='done' else None
            obj.save(); form.save_m2m()
            if items is not None: items.instance=obj; items.save()
            if kind=='customers':
                if previous_stage!=obj.stage:
                    JourneyEvent.objects.create(customer=obj,from_stage=previous_stage,to_stage=obj.stage,actor=request.user,note='Customer profile updated' if pk else 'Customer created')
                if pk and old_owner!=obj.owner_id:
                    obj.orders.update(owner=obj.owner); obj.followups.update(owner=obj.owner)
                    audit(request.user,'reassign-related',obj)
            audit(request.user,'update' if pk else 'create',obj)
            messages.success(request,f'{singular} saved.')
            return redirect('detail',kind=kind,pk=obj.pk) if kind in ('orders','customers') else redirect('listing',kind=kind)
    return render(request,'crm/form.html',{'form':form,'items':items,'kind':kind,'title':('Edit ' if pk else 'New ')+singular.lower(),'revision':old_updated,'products':list(Product.objects.filter(active=True).values('id','name','price')) if items is not None else []})

@login_required
@transaction.atomic
def delete(request,kind,pk):
    if kind in ('products','targets') and not request.user.is_staff: raise PermissionDenied
    obj=get_object_or_404(scoped(request.user,kind).select_for_update(),pk=pk)
    if request.method=='POST':
        try:
            with transaction.atomic():
                audit(request.user,'delete',obj); obj.delete()
        except ProtectedError:
            messages.error(request,'This record has linked orders, products or payments. Remove dependent records first, or keep it for reporting.')
        else:
            messages.success(request,'Record deleted.'); return redirect('listing',kind=kind)
    return render(request,'crm/delete.html',{'object':obj,'kind':kind,'title':'Delete record'})

@login_required
def detail(request,kind,pk):
    if kind not in ('customers','orders'): raise Http404
    obj=get_object_or_404(scoped(request.user,kind),pk=pk)
    context={'object':obj,'kind':kind,'title':str(obj)}
    if kind=='customers':
        context.update(orders=obj.orders.prefetch_related('items','payments'),followups=obj.followups.all(),journey=obj.journey.select_related('actor'),stages=Customer.STAGES)
    return render(request,'crm/detail.html',context)

@login_required
def dashboard(request):
    today=timezone.localdate(); first=today.replace(day=1)
    try:
        start=date.fromisoformat(request.GET.get('start') or first.isoformat())
        end=date.fromisoformat(request.GET.get('end') or today.isoformat())
        if start>end: raise ValueError
    except ValueError: return HttpResponse('Choose a valid date range with start before end.',status=400)
    allorders=list(scoped(request.user,'orders'))
    orders=[o for o in allorders if start<=o.order_date<=end and o.status in REVENUE_STATUSES]
    customers=scoped(request.user,'customers')
    followups=scoped(request.user,'followups')
    revenue=sum((o.total for o in orders),Decimal(0))
    collected=sum((p.amount for p in scoped(request.user,'payments').filter(paid_on__range=[start,end])),Decimal(0))
    outstanding=sum((o.balance for o in allorders if o.status in REVENUE_STATUSES),Decimal(0))
    pipeline=list(customers.values_list('stage',flat=True))
    funnel=[{'label':label,'value':pipeline.count(key)} for key,label in Customer.STAGES]
    # Fixed six-month graph ending in the selected report's end month.
    months=[]
    for offset in range(5,-1,-1):
        n=end.year*12+end.month-1-offset; y,m=divmod(n,12)
        total=sum((o.total for o in allorders if o.order_date.year==y and o.order_date.month==m+1 and o.status in REVENUE_STATUSES),Decimal(0))
        months.append({'label':date(y,m+1,1).strftime('%b %y'),'value':float(total)})
    staff=[]
    users=get_user_model().objects.filter(is_active=True) if request.user.is_staff else [request.user]
    for user in users:
        own=[o for o in orders if o.owner_id==user.pk]
        actual=sum((o.total for o in own),Decimal(0))
        target=SalesTarget.objects.filter(user=user,month=end.replace(day=1)).first()
        month_actual=sum((o.total for o in allorders if o.owner_id==user.pk and o.order_date.year==end.year and o.order_date.month==end.month and o.status in REVENUE_STATUSES),Decimal(0))
        staff.append({'user':user,'sales':actual,'count':len(own),'clients':customers.filter(owner=user,stage='client').count(),'overdue':followups.filter(owner=user,status='open',due_at__lt=timezone.now()).count(),'target':target.amount if target else None,'achievement':round(month_actual/target.amount*100) if target else 0,'bar':min(100,round(month_actual/target.amount*100)) if target else 0,'month_actual':month_actual})
    context={'title':'Sales overview','start':start,'end':end,'revenue':revenue,'collected':collected,'outstanding':outstanding,'customer_count':customers.count(),'overdue':followups.filter(status='open',due_at__lt=timezone.now()).count(),'due':followups.filter(status='open')[:6],'recent':orders[:6],'funnel':funnel,'months':months,'staff':staff,'operations':request.user.is_staff,'logs':AuditLog.objects.select_related('actor')[:15] if request.user.is_staff else [],'month_label':end.strftime('%B %Y')}
    return render(request,'crm/dashboard.html',context)

@login_required
def operations(request):
    if not request.user.is_staff: raise PermissionDenied
    return dashboard(request)

def safe_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r','\n')): return "'"+value
    return value

@login_required
def export(request,kind,fmt):
    if fmt not in ('csv','xlsx'): raise Http404
    try: qs=filtered(request,kind)
    except ValueError as exc: return HttpResponse(str(exc),status=400)
    headers={
      'customers':['ID','Name','Company','Email','Phone','Stage','Owner','Estimated value (RM)'],
      'orders':['Order','Customer','Owner','Date','Status','Total (RM)','Paid (RM)','Balance (RM)'],
      'followups':['ID','Customer','Owner','Subject','Channel','Due (Malaysia)','Status','Priority','Notes'],
      'products':['ID','SKU','Name','Price (RM)','Active'],
      'payments':['ID','Order','Customer','Amount (RM)','Paid on','Method','Reference'],
      'targets':['ID','User','Month','Target (RM)'],
    }[kind]
    def row(o):
        if kind=='customers': return [o.pk,o.name,o.company,o.email,o.phone,o.get_stage_display(),o.owner.username,o.estimated_value]
        if kind=='orders': return [str(o),o.customer.name,o.owner.username,o.order_date.isoformat(),o.get_status_display(),o.total,o.paid,o.balance]
        if kind=='followups': return [o.pk,o.customer.name,o.owner.username,o.subject,o.get_kind_display(),timezone.localtime(o.due_at).strftime('%Y-%m-%d %H:%M'),o.get_status_display(),o.priority,o.notes]
        if kind=='products': return [o.pk,o.sku,o.name,o.price,o.active]
        if kind=='payments': return [o.pk,str(o.order),o.order.customer.name,o.amount,o.paid_on.isoformat(),o.get_method_display(),o.reference]
        return [o.pk,o.user.username,o.month.isoformat(),o.amount]
    AuditLog.objects.create(actor=request.user,action='export-'+fmt,entity=kind,object_id=0,label=f'Filtered {kind} export')
    if fmt=='csv':
        response=HttpResponse(content_type='text/csv; charset=utf-8'); response.write('\ufeff')
        writer=csv.writer(response); writer.writerow(headers)
        for o in qs: writer.writerow([safe_cell(v) for v in row(o)])
    else:
        wb=Workbook(); ws=wb.active; ws.title=kind.title(); ws.append(headers)
        for o in qs: ws.append([safe_cell(v) for v in row(o)])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=24
        from openpyxl.styles import Font, PatternFill
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='154F44')
        buffer=BytesIO(); wb.save(buffer)
        response=HttpResponse(buffer.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=f'attachment; filename="{kind}-{timezone.localdate()}.{fmt}"'
    return response
